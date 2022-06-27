from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference

wb = Workbook(write_only=True)
sheet = wb.create_sheet()


rows =[
    ('类别','销售A组','销售B组'),
    ('手机',40,30),
    ('平板电脑',50,60),
    ('笔记本', 80, 70),
    ('VR眼镜', 20, 10),
]

# 向表单中添加行

for row in rows:
    sheet.append(row)

# 创建图标对象
chart = BarChart()
chart.type = 'col'
chart.style = 10
# 设置图标标题
chart.title = "销售统计图"

# 设置图表纵轴的标题
chart.y_axis.title = "销量"

# 设置图标横轴的标题
chart.x_axis.title = "商品类别"

# 设置数据的范围
data = Reference(sheet,min_col=2,min_row=1,max_row=5,max_col=3)

# 设置分类的范围
cats = Reference(sheet,min_col=1,min_row=2,max_row=5)

# 给图表添加数据
chart.add_data(data,titles_from_data=True)

# 给图表设置分类
chart.set_categories(cats)
chart.shpe = 4

# 将图表插入到工作表的指定单元格
sheet.add_chart(chart,"A10")

wb.save("./data/销量统计表.xlsx")



